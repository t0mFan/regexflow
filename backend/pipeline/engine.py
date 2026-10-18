import csv, json, os, shutil, uuid
from pathlib import Path
from functools import partial
from pyspark.sql import SparkSession, functions as F, types as T

def spark_session():
    from django.conf import settings
    session = (SparkSession.builder.appName("RegexFlow").master(settings.SPARK_MASTER)
        .config("spark.sql.shuffle.partitions", settings.SPARK_PARTITIONS)
        .config("spark.sql.adaptive.enabled", "true")
        .config("spark.ui.enabled", "false")
        .config("spark.driver.bindAddress", os.getenv("SPARK_DRIVER_BIND", "127.0.0.1"))
        .config("spark.driver.host", os.getenv("SPARK_DRIVER_HOST", "127.0.0.1"))
        .config("spark.executor.cores", "1")
        .getOrCreate())
    session.sparkContext.setLogLevel("WARN")
    return session

def excel_to_csv(source, target):
    from openpyxl import load_workbook
    wb = load_workbook(source, read_only=True, data_only=True)
    try:
        with open(target, "w", newline="", encoding="utf-8") as out:
            writer = csv.writer(out)
            for row in wb.active.iter_rows(values_only=True):
                writer.writerow(["" if x is None else str(x) for x in row])
    finally:
        wb.close()

def load_frame(spark, path, partitions):
    if str(path).lower().endswith(".xlsx"):
        target = str(path) + ".csv"
        excel_to_csv(path, target)
        path = target
    with open(path, newline="", encoding="utf-8-sig") as stream:
        headers = next(csv.reader(stream), [])
    if not headers or any(not h.strip() for h in headers) or len(headers) != len(set(headers)):
        raise ValueError("File requires non-empty, unique column headers.")
    frame = spark.read.option("header", True).option("inferSchema", False).option("multiLine", True).option("mode", "FAILFAST").csv(str(path))
    if not frame.columns or len(set(frame.columns)) != len(frame.columns):
        raise ValueError("File requires unique column headers.")
    # CSV multiline parsing can be a single input task. Repartition once after
    # ingestion, then all regex work and page generation run across partitions.
    return frame.repartition(partitions)

def transform_partition(rows, columns, operation, plan, replacement):
    import re2
    pattern = re2.compile(plan["pattern"]) if operation != "normalize" else None
    for row in rows:
        values = row.asDict()
        for col in columns:
            value = values[col]
            if value is None:
                continue
            if len(value) > 1_000_000:
                raise ValueError("A cell exceeds the one-million-character limit.")
            if operation == "replace":
                # Callable replacement treats dollar signs and backslashes literally.
                values[col] = pattern.sub(lambda match: replacement, value)
            elif operation == "extract":
                match = pattern.search(value)
                values[col] = match.group(0) if match else None
            else:
                action = plan["normalization"]
                values[col] = {"trim": lambda: value.strip(), "lower": lambda: value.strip().lower(), "upper": lambda: value.strip().upper(), "collapse_spaces": lambda: " ".join(value.split())}[action]()
        yield values

def transform(frame, columns, operation, plan, replacement):
    fn = partial(transform_partition, columns=columns, operation=operation, plan=plan, replacement=replacement)
    return frame.sparkSession.createDataFrame(frame.rdd.mapPartitions(fn), frame.schema)

def write_partition(index, rows, directory, page_size=100):
    # Only a single browser page is ever buffered. Spark retries use unique
    # filenames, and only the successful attempt's manifest is published.
    attempt = uuid.uuid4().hex
    names, batch, total = [], [], 0
    def flush():
        name = f"p{index}-{attempt}-{len(names):07d}.json"
        Path(directory, name).write_text(json.dumps(batch, ensure_ascii=False), encoding="utf-8")
        names.append(name)
    for row in rows:
        batch.append(row.asDict())
        total += 1
        if len(batch) == page_size:
            flush()
            batch = []
    if batch:
        flush()
    yield {"partition": index, "pages": names, "rows": total}

def materialize(frame, directory):
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    shards = frame.rdd.mapPartitionsWithIndex(partial(write_partition, directory=str(directory))).collect()
    shards.sort(key=lambda item: item["partition"])
    manifest = {"columns": frame.columns, "pages": [name for shard in shards for name in shard["pages"]], "rows": sum(shard["rows"] for shard in shards), "page_size": 100}
    tmp = directory / "manifest.tmp"
    tmp.write_text(json.dumps(manifest), encoding="utf-8")
    os.replace(tmp, directory / "manifest.json")
    return {"columns": manifest["columns"], "rows": manifest["rows"], "page_count": len(manifest["pages"]), "page_size": 100}


def count_matches_partition(rows, columns, operation, plan):
    import re2
    pattern = re2.compile(plan["pattern"]) if operation != "normalize" else None
    populated, matched = 0, 0
    for row in rows:
        values = row.asDict()
        for name in columns:
            value = values[name]
            if value is None or value == "":
                continue
            populated += 1
            if pattern is None or pattern.search(value) is not None:
                matched += 1
    yield (populated, matched)


def matching_stats(frame, columns, operation, plan):
    fn = partial(count_matches_partition, columns=columns, operation=operation, plan=plan)
    populated, matched = frame.rdd.mapPartitions(fn).fold((0, 0), lambda a, b: (a[0]+b[0], a[1]+b[1]))
    if operation == "extract" and populated and not matched:
        raise ValueError("The rule matched no non-empty values in the selected columns. Extraction was stopped to avoid turning the entire column into nulls. Specify what to extract and check the target column.")
    warnings = []
    if operation == "replace" and populated and not matched:
        warnings.append("Nothing matched, so this replacement did not change the data. Check the rule and target columns.")
    return {"populated_cells": populated, "matched_cells": matched, "warnings": warnings}
